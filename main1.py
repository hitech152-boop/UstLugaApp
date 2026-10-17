import io
import os
import re
import threading
import urllib.request
import ssl
import openpyxl
from openpyxl.utils import get_column_letter

from kivy.app import App
from kivy.lang import Builder
from kivy.uix.boxlayout import BoxLayout
from kivy.uix.label import Label
from kivy.uix.button import Button
from kivy.uix.scrollview import ScrollView
from kivy.uix.popup import Popup
from kivy.uix.textinput import TextInput
from kivy.properties import StringProperty, NumericProperty, ListProperty, BooleanProperty
from kivy.clock import Clock
from kivy.metrics import dp

# ---------- НАСТРОЙКИ ----------
FILE_ID = '1HQPiYxLhi6V56VSZYJUYZ1vL3SfdeNWf'
SHEET_NAME = 'Анализ по объектам'

# Индексы столбцов (0-based): D=3, E=4, F=5, G=6, M=12
USE_COLS = [3, 4, 5, 6, 12]

COL_WIDTHS = [dp(320), dp(220), dp(65), dp(65), dp(65)]
N_COLS = len(USE_COLS)
TOTAL_WIDTH = sum(COL_WIDTHS)


# ---------- KV-РАЗМЕТКА ----------
KV = '''
<CellLabel>:
    size_hint: None, None
    height: dp(55) 
    font_size: '12sp'
    color: 0.05, 0.05, 0.05, 1
    text_size: self.width, None
    halign: 'left'
    valign: 'middle'
    padding_x: dp(4)
    canvas.before:
        Color:
            rgba: (0.95, 0.97, 1, 1) if self.row_index % 2 == 0 else (1, 1, 1, 1)
        Rectangle:
            pos: self.pos
            size: self.size
    canvas.after:
        Color:
            rgba: 0.7, 0.7, 0.7, 1
        Line:
            points: self.right - 1, self.y, self.right - 1, self.top
            width: 1
        Color:
            rgba: 0.85, 0.85, 0.85, 1
        Line:
            points: self.x, self.y, self.right, self.y
            width: 1

<CalcButton>:
    size_hint: None, None
    height: dp(55)
    background_normal: ''
    background_color: (0, 0, 0, 0)
    color: 0.1, 0.3, 0.8, 1
    font_size: '12sp'
    text_size: self.width, None
    halign: 'left'
    valign: 'middle'
    padding_x: dp(4)
    canvas.before:
        Color:
            rgba: (0.95, 0.97, 1, 1) if self.row_index % 2 == 0 else (1, 1, 1, 1)
        Rectangle:
            pos: self.pos
            size: self.size
    canvas.after:
        Color:
            rgba: 0.7, 0.7, 0.7, 1
        Line:
            points: self.right - 1, self.y, self.right - 1, self.top
            width: 1
        Color:
            rgba: 0.85, 0.85, 0.85, 1
        Line:
            points: self.x, self.y, self.right, self.y
            width: 1

<HeaderCell>:
    size_hint: None, None
    height: dp(50)
    font_size: '12sp'
    bold: True
    color: 1, 1, 1, 1
    text_size: self.width, None
    halign: 'left'
    valign: 'middle'
    padding_x: dp(4)
    canvas.before:
        Color:
            rgba: 0.18, 0.35, 0.55, 1
        Rectangle:
            pos: self.pos
            size: self.size
    canvas.after:
        Color:
            rgba: 0.1, 0.2, 0.3, 1
        Line:
            points: self.right - 1, self.y, self.right - 1, self.top
            width: 1
        Color:
            rgba: 0.1, 0.2, 0.3, 1
        Line:
            points: self.x, self.y, self.right, self.y
            width: 1

<TableRow>:
    orientation: 'horizontal'
    size_hint: None, None
    height: dp(55) 
    width: app.total_width

<HeaderRow>:
    orientation: 'horizontal'
    size_hint: None, None
    height: dp(50)
    width: app.total_width

<MainScreen>:
    orientation: 'vertical'
    padding: [dp(24), dp(8), dp(16), dp(8)]
    spacing: dp(6)

    BoxLayout:
        size_hint_y: None
        height: dp(44)
        spacing: dp(6)

        Button:
            text: 'Обновить'
            size_hint_x: None
            width: dp(100)
            font_size: '13sp'
            on_release: root.load_data()

        TextInput:
            id: search_input
            hint_text: 'Введите тег, например L1-41D02'
            multiline: False
            font_size: '13sp'
            on_text_validate: root.do_search()

        Button:
            text: 'Найти'
            size_hint_x: None
            width: dp(80)
            font_size: '13sp'
            on_release: root.do_search()

        Button:
            text: 'Все'
            size_hint_x: None
            width: dp(70)
            font_size: '13sp'
            on_release: root.show_all()

        Button:
            text: 'Теги'
            size_hint_x: None
            width: dp(70)
            font_size: '13sp'
            on_release: root.show_tags()

    Label:
        text: 'by KarpyukDA'
        size_hint_y: None
        height: dp(16)
        font_size: '10sp'
        color: 0.5, 0.5, 0.5, 1
        text_size: self.width, None
        halign: 'right'
        valign: 'middle'

    Label:
        text: root.status_text
        size_hint_y: None
        height: dp(28)
        font_size: '12sp'
        color: 0.1, 0.1, 0.1, 1
        text_size: self.width, None
        halign: 'left'
        valign: 'middle'

    RecycleView:
        id: header_rv
        size_hint_y: None
        height: dp(50)
        do_scroll_x: True
        do_scroll_y: False
        bar_width: 0
        viewclass: 'HeaderRow'
        RecycleBoxLayout:
            orientation: 'vertical'
            size_hint: None, None
            width: app.total_width
            height: dp(50)
            default_size: None, dp(50)

    RecycleView:
        id: rv
        do_scroll_x: True
        do_scroll_y: True
        bar_width: dp(24) 
        bar_margin: dp(4)
        scroll_type: ['bars', 'content']
        bar_color: 0.4, 0.4, 0.4, 1
        bar_inactive_color: 0.6, 0.6, 0.6, 1
        on_scroll_x: header_rv.scroll_x = self.scroll_x
        viewclass: 'TableRow'
        RecycleBoxLayout:
            orientation: 'vertical'
            size_hint: None, None
            width: app.total_width
            height: self.minimum_height
            default_size: None, dp(55)
'''

Builder.load_string(KV)


# ---------- ВСПОМОГАТЕЛЬНЫЕ КЛАССЫ ----------

class LoadingPopup(Popup):
    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self.title = 'Подождите'
        self.size_hint = (0.6, 0.25)
        self.auto_dismiss = False
        
        self.lbl = Label(text='Загрузка данных', font_size='16sp')
        self.content = self.lbl
        self._dots = 0
        self._event = Clock.schedule_interval(self._animate, 0.4)

    def _animate(self, dt):
        self._dots = (self._dots + 1) % 4
        self.lbl.text = 'Загрузка данных' + '.' * self._dots

    def on_dismiss(self):
        Clock.unschedule(self._event)


class CalcPopup(Popup):
    def __init__(self, rkc_value, **kwargs):
        super().__init__(**kwargs)
        self.title = 'Калькулятор'
        self.size_hint = (0.7, 0.5)
        self.auto_dismiss = True
        self.rkc_val = rkc_value
        
        layout = BoxLayout(orientation='vertical', padding=dp(15), spacing=dp(10))
        
        info_text = f'Объем × {rkc_value:.2f}'
        layout.add_widget(Label(
            text=info_text, 
            size_hint_y=None, 
            height=dp(30), 
            font_size='16sp',
            halign='center'
        ))
        
        layout.add_widget(Label(
            text='Введите объем:', 
            size_hint_y=None, 
            height=dp(25), 
            halign='left',
            font_size='14sp'
        ))
        
        self.input = TextInput(
            hint_text='0.00', 
            multiline=False, 
            font_size='22sp',
            input_filter='float'
        )
        self.input.bind(text=self.on_input)
        layout.add_widget(self.input)
        
        self.result_lbl = Label(
            text='Итого: 0.00 часов', 
            font_size='22sp', 
            bold=True, 
            color=(0.1, 0.5, 0.1, 1)
        )
        layout.add_widget(self.result_lbl)
        
        close_btn = Button(text='Закрыть', size_hint_y=None, height=dp(45), font_size='16sp')
        close_btn.bind(on_release=self.dismiss)
        layout.add_widget(close_btn)
        
        self.content = layout

    def on_input(self, instance, value):
        val_str = value.replace(',', '.')
        try:
            vol = float(val_str)
            res = vol * self.rkc_val
            self.result_lbl.text = f'Итого: {res:.2f} часов'
        except ValueError:
            self.result_lbl.text = 'Итого: 0.00 часов'


class CalcButton(Button):
    row_index = NumericProperty(0)
    rkc_value = NumericProperty(0)

    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self.bind(on_release=self.do_calc)

    def do_calc(self, instance):
        if self.rkc_value > 0:
            CalcPopup(rkc_value=self.rkc_value).open()


class CellLabel(Label):
    row_index = NumericProperty(0)


class HeaderCell(Label):
    pass


class TableRow(BoxLayout):
    row_index = NumericProperty(0)
    row_data = ListProperty([])
    is_bold = BooleanProperty(False)

    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self.cells = []
        for i, w in enumerate(COL_WIDTHS):
            if i == 4:
                c = CalcButton()
            else:
                c = CellLabel()
                if i == 0:
                    c.padding_x = dp(16)
                    
            c.size_hint_x = None
            c.width = w
            self.add_widget(c)
            self.cells.append(c)

    def on_row_index(self, instance, value):
        for c in self.cells:
            c.row_index = value

    def on_is_bold(self, instance, value):
        for c in self.cells:
            if hasattr(c, 'bold'):
                c.bold = value

    def on_row_data(self, instance, value):
        for i, val in enumerate(value):
            if i < len(self.cells):
                self.cells[i].text = str(val)
                if i == 4:
                    try:
                        clean = str(val).replace(' ', '').replace('\xa0', '').replace(',', '.')
                        self.cells[i].rkc_value = float(clean)
                    except ValueError:
                        self.cells[i].rkc_value = 0


class HeaderRow(BoxLayout):
    row_data = ListProperty([])

    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self.cells = []
        for i, w in enumerate(COL_WIDTHS):
            c = HeaderCell()
            c.size_hint_x = None
            c.width = w
            if i == 0:
                c.padding_x = dp(16)
            self.add_widget(c)
            self.cells.append(c)

    def on_row_data(self, instance, value):
        for i, val in enumerate(value):
            if i < len(self.cells):
                self.cells[i].text = str(val)


class MainScreen(BoxLayout):
    status_text = StringProperty('Приложение запущено')
    tags = ListProperty([])

    data_rows = []
    header_row_idx = None
    tags_by_group = {}
    _tags_popup = None
    loading_popup = None

    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        Clock.schedule_once(lambda dt: self.load_data(), 0.5)

    def load_data(self):
        self.status_text = 'Скачиваем файл из Google Drive...'
        self.loading_popup = LoadingPopup()
        self.loading_popup.open()
        
        threading.Thread(target=self._thread_work, daemon=True).start()

    def _thread_work(self):
        try:
            file_obj = self._download_file()
            if not file_obj:
                self._finish_loading('Ошибка: не удалось скачать файл')
                return

            data_rows = self._read_excel_openpyxl(file_obj)
            if data_rows is None:
                self._finish_loading('Ошибка: данные пусты')
                return

            Clock.schedule_once(lambda dt: self._process_data(data_rows), 0)
            
        except Exception as e:
            print(f"Ошибка в потоке: {e}")
            Clock.schedule_once(lambda dt: self._finish_loading(f'Ошибка: {str(e)[:80]}'), 0)

    def _finish_loading(self, status):
        if self.loading_popup:
            self.loading_popup.dismiss()
            self.loading_popup = None
        self.status_text = status

    def _process_data(self, data_rows):
        self.data_rows = data_rows
        
        self.header_row_idx = None
        for i, row in enumerate(self.data_rows):
            if i >= 20: break
            val = row[0] if row else ''
            if str(val).strip() == 'Вид работ':
                self.header_row_idx = i
                break

        self._collect_tags()
        self.show_all()
        
        if self.loading_popup:
            self.loading_popup.dismiss()
            self.loading_popup = None
            
        self.status_text = f'Загружено строк: {len(self.data_rows)}, найдено тегов: {len(self.tags)}'

    def _download_file(self):
        url = f'https://drive.google.com/uc?export=download&id={FILE_ID}'
        
        # Создаем небезопасный контекст SSL (для обхода корпоративных сертификатов)
        context = ssl._create_unverified_context()
        
        try:
            req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
            response = urllib.request.urlopen(req, context=context)
            data = response.read()
            
            # Проверка, что скачался не HTML
            if data.startswith(b'<!DOCTYPE html>') or data.startswith(b'<html'):
                print("ОШИБКА: Файл не является публичным. Скачалась HTML-страница.")
                return None
                
            file = io.BytesIO(data)
            file.seek(0)
            return file
        except Exception as e:
            print(f"Ошибка скачивания: {e}")
            return None

    def _read_excel_openpyxl(self, file_obj):
        try:
            wb = openpyxl.load_workbook(file_obj, data_only=True, read_only=True)
            if SHEET_NAME not in wb.sheetnames:
                return None
                
            ws = wb[SHEET_NAME]
            all_rows = []
            
            for row in ws.iter_rows(values_only=True):
                vals = []
                for idx in USE_COLS:
                    if idx < len(row):
                        vals.append(row[idx])
                    else:
                        vals.append(None)
                all_rows.append(vals)
            
            wb.close()
            return all_rows
        except Exception as e:
            print(f"Ошибка чтения openpyxl: {e}")
            return None

    def _collect_tags(self):
        self.tags_by_group = {}
        all_tags = set()
        
        for row in self.data_rows:
            if not row: continue
            d = str(row[0]).strip() if row[0] is not None else ''
            if re.match(r'^L\d+-\d+', d) and d.lower() != 'nan':
                all_tags.add(d)
                match = re.match(r'^(L\d+)', d)
                if match:
                    prefix = match.group(1)
                    if prefix not in self.tags_by_group:
                        self.tags_by_group[prefix] = set()
                    self.tags_by_group[prefix].add(d)
        
        for k in self.tags_by_group:
            self.tags_by_group[k] = sorted(self.tags_by_group[k])
            
        self.tags = sorted(all_tags)

    def _find_all_blocks(self, tag):
        if not self.data_rows:
            return None
            
        start_indices = []
        for i, row in enumerate(self.data_rows):
            if row and str(row[0]).strip() == 'ВелесстройМонтаж':
                start_indices.append(i)
        
        matching_blocks = []
        
        for idx, start_idx in enumerate(start_indices):
            end_idx = start_indices[idx + 1] if idx + 1 < len(start_indices) else len(self.data_rows)
            block = self.data_rows[start_idx:end_idx]
            
            found = False
            for row in block:
                if row and str(row[0]).strip() == tag:
                    found = True
                    break
            
            if found:
                matching_blocks.append(block)
        
        if not matching_blocks:
            return None

        result = []
        for i, b in enumerate(matching_blocks):
            if i > 0:
                result.append([None]*len(b[0]) if b else [None]*5)
            result.extend(b)
            
        return result

    def do_search(self):
        tag = self.ids.search_input.text.strip()
        if not tag:
            self.status_text = 'Введите тег для поиска'
            return
        if not self.data_rows:
            self.status_text = 'Сначала загрузите данные'
            return

        block_rows = self._find_all_blocks(tag)
        if block_rows is None:
            self.status_text = f'Блок с тегом «{tag}» не найден'
            return

        header_vals = ['D', 'E', 'F', 'G', 'M']
        for row in block_rows:
            if row and str(row[0]).strip() == 'Вид работ':
                header_vals = [str(v) for v in row]
                break

        self._set_header(header_vals)
        self._display_rows(block_rows)
        
        count_blocks = sum(1 for row in block_rows if row and str(row[0]).strip() == 'ВелесстройМонтаж')
        self.status_text = f'Найдено блоков с тегом «{tag}»: {count_blocks}'

    def show_all(self):
        if not self.data_rows:
            return
        if self.header_row_idx is not None:
            header_vals = [str(v) for v in self.data_rows[self.header_row_idx]]
            self._set_header(header_vals)
            data_rows = self.data_rows[self.header_row_idx + 1:]
        else:
            self._set_header(['D', 'E', 'F', 'G', 'M'])
            data_rows = self.data_rows

        self._display_rows(data_rows)
        self.status_text = f'Показаны все данные'

    def show_tags(self):
        if not self.tags_by_group:
            self.status_text = 'Список тегов пуст. Сначала загрузите данные.'
            return

        content = BoxLayout(orientation='vertical', padding=dp(8), spacing=dp(6))
        scroll = ScrollView()
        inner = BoxLayout(orientation='vertical', size_hint_y=None, spacing=dp(4))
        inner.bind(minimum_height=inner.setter('height'))

        groups = sorted(self.tags_by_group.keys())
        
        for group in groups:
            count = len(self.tags_by_group[group])
            btn = Button(
                text=f'Группа {group}  (тегов: {count})',
                size_hint_y=None,
                height=dp(44),
                font_size='14sp',
                background_color=(0.2, 0.4, 0.6, 1)
            )
            btn.bind(on_release=lambda x, g=group: self._show_group_tags(g))
            inner.add_widget(btn)

        scroll.add_widget(inner)
        content.add_widget(scroll)

        close_btn = Button(text='Закрыть', size_hint_y=None, height=dp(42))
        content.add_widget(close_btn)

        popup = Popup(
            title='Выберите группу тегов',
            content=content,
            size_hint=(0.7, 0.8),
        )
        close_btn.bind(on_release=popup.dismiss)
        self._tags_popup = popup
        popup.open()

    def _show_group_tags(self, group):
        if self._tags_popup:
            self._tags_popup.dismiss()
            
        tags = self.tags_by_group.get(group, [])
        
        content = BoxLayout(orientation='vertical', padding=dp(8), spacing=dp(4))
        scroll = ScrollView()
        inner = BoxLayout(orientation='vertical', size_hint_y=None, spacing=dp(2))
        inner.bind(minimum_height=inner.setter('height'))

        for tag in tags:
            btn = Button(
                text=tag,
                size_hint_y=None,
                height=dp(36),
                font_size='13sp',
            )
            btn.bind(on_release=lambda x, t=tag: self._select_tag(t))
            inner.add_widget(btn)

        scroll.add_widget(inner)
        content.add_widget(scroll)

        close_btn = Button(text='Назад', size_hint_y=None, height=dp(42))
        content.add_widget(close_btn)
        
        popup = Popup(
            title=f'Теги группы {group} ({len(tags)})',
            content=content,
            size_hint=(0.7, 0.8),
        )
        
        def go_back(instance):
            popup.dismiss()
            self.show_tags()
            
        close_btn.bind(on_release=go_back)
        self._tags_popup = popup
        popup.open()

    def _select_tag(self, tag):
        self.ids.search_input.text = tag
        if self._tags_popup:
            self._tags_popup.dismiss()
        self.do_search()

    def _fmt(self, v):
        if v is None:
            return ''
        if isinstance(v, bool):
            return ''
        if isinstance(v, float):
            if v == int(v):
                return str(int(v))
            return f'{v:.2f}'
        return str(v)

    def _set_header(self, values):
        vals = [str(v) for v in values]
        if len(vals) > 2:
            vals[2] = 'Выполнено'
        if len(vals) > 4:
            vals[4] = 'РКЦ (нажми)'
            
        self.ids.header_rv.data = [{'row_data': vals}]

    def _clean_text(self, text):
        if text is None:
            return ''
        return str(text).replace('\xa0', ' ').replace('\n', ' ').replace('\r', '').strip()

    def _display_rows(self, rows_data):
        rows = []
        display_index = 0
        
        for row_vals in rows_data:
            first_col_val = self._clean_text(row_vals[0])
            
            if not first_col_val or first_col_val.lower() == 'nan':
                continue
            if first_col_val in ['Вид работ', 'ИТОГО']:
                continue
                
            is_bold = False
            if first_col_val == 'ВелесстройМонтаж' or re.match(r'^L\d+-\d+', first_col_val):
                is_bold = True
            else:
                has_numbers = False
                for idx in [2, 3, 4]:
                    if idx < len(row_vals):
                        val = row_vals[idx]
                        if val is not None and not isinstance(val, bool):
                            try:
                                float(val)
                                has_numbers = True
                                break
                            except (ValueError, TypeError):
                                pass
                
                if not has_numbers:
                    is_bold = True
                    
            formatted = [self._fmt(v) for v in row_vals]
                
            rows.append({
                'row_data': formatted, 
                'row_index': display_index, 
                'is_bold': is_bold
            })
            display_index += 1
            
        self.ids.rv.data = rows


class UstLugaApp(App):
    total_width = NumericProperty(TOTAL_WIDTH)

    def build(self):
        self.title = 'Трудозатраты Усть-Луга'
        return MainScreen()


if __name__ == '__main__':
    UstLugaApp().run()